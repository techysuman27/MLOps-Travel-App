"""Create the separate course workflow report from the maintained narrative."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

root = Path(__file__).resolve().parents[1]
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.75)
normal = doc.styles['Normal']
normal.font.name, normal.font.size = 'Calibri', Pt(11)
normal.paragraph_format.space_after = Pt(8)
normal.paragraph_format.line_spacing = 1.12
for style in ['Title', 'Heading 1', 'Heading 2']:
    doc.styles[style].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Title'].font.size = Pt(26)
doc.styles['Heading 1'].font.size = Pt(17)
doc.core_properties.author = 'techysuman27'
doc.core_properties.title = 'Voyage Analytics MLOps Workflow Report'
doc.add_paragraph('Voyage Analytics MLOps Workflow Report', 'Title')
doc.add_paragraph('AlmaBetter Productionization of ML Systems\nIndividual capstone by techysuman27')
doc.add_paragraph('Repository: https://github.com/techysuman27/MLOps-Travel-App\nPublic app: https://voyage-analytics-suman.streamlit.app/\nPublic API: https://voyage-analytics-api-kqju.onrender.com')
doc.add_paragraph('This report explains how the travel models move from validated historical data to a public prediction service. It records the serving, packaging, scaling, scheduling, integration and tracking workflow, supported by captured execution evidence. It also identifies checks that are still pending and limits on interpreting the results.')

for block in (root / 'docs/workflow.md').read_text(encoding='utf-8').split('\n\n'):
    block = block.strip()
    if not block or block.startswith('# ') or block.startswith('**Author:') or block.startswith('This is the source narrative'):
        continue
    if block.startswith('## '):
        heading = re.sub(r'[^\w\s]', '', block[3:])
        doc.add_heading(heading, level=1)
    elif block.startswith('!['):
        match = re.match(r'!\[(.*?)\]\((.*?)\)', block)
        if match:
            image = root / 'docs' / match.group(2)
            if image.exists():
                with Image.open(image) as im:
                    width = min(7.0, 6.6 * im.width / im.height)
                p = doc.add_paragraph()
                p.paragraph_format.keep_with_next = True
                shape = p.add_run().add_picture(str(image), width=Inches(width))
                shape._inline.docPr.set('descr', match.group(1))
                doc.add_paragraph(match.group(1), 'Caption')
    else:
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', block).replace('`', '')
        doc.add_paragraph(text)
output = root / 'submission/Voyage_Analytics_Workflow.docx'
output.parent.mkdir(exist_ok=True)
doc.save(output)
print(output)
